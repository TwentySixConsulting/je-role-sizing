import docx, json, os, re, sys
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
import openpyxl

# ---------- 1. parse the scheme docx (authoritative criteria + points) ----------
# Source documents: pass a directory as the first argument, otherwise look in
# ~/Downloads, where the originals live.
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser("~/Downloads")
SCHEME_DOC = os.path.join(SRC, "Our role sizing scheme v2.docx")
WORKBOOK = os.path.join(SRC, "JE Scheme Database 2022.xlsx")
for f in (SCHEME_DOC, WORKBOOK):
    if not os.path.exists(f):
        sys.exit(f"Missing source document: {f}\nUsage: generate-scheme.py [directory holding the two source files]")

d = docx.Document(SCHEME_DOC)
body = d.element.body
cornerstones = []   # {name, intro}
factors = []        # {name, cornerstone, levels:[{level, description, points}]}
cur_cs = None
pending_factor = None
last_para = None

def cell_texts(row):
    out, seen = [], set()
    for c in row.cells:
        if id(c._tc) in seen: continue
        seen.add(id(c._tc))
        out.append(" ".join(x.strip() for x in c.text.split("\n") if x.strip()))
    return out

for child in body.iterchildren():
    if child.tag == qn('w:p'):
        t = Paragraph(child, d).text.strip()
        if not t: continue
        m = re.match(r'Cornerstone\s+(\d+):\s*(.+)', t)
        if m:
            cur_cs = {"num": int(m.group(1)), "name": m.group(2).strip(), "intro": ""}
            cornerstones.append(cur_cs); last_para = None; continue
        m = re.match(r'Factor\s+(\d+):\s*(.+)', t)
        if m:
            pending_factor = {"num": int(m.group(1)), "name": m.group(2).strip(),
                              "cornerstone": cur_cs["name"], "levels": []}
            continue
        if cur_cs and not cur_cs["intro"] and pending_factor is None:
            cur_cs["intro"] = re.sub(r'\s+', ' ', t)
    elif child.tag == qn('w:tbl'):
        tb = Table(child, d)
        rows = [cell_texts(r) for r in tb.rows]
        assert pending_factor, "table without factor"
        for r in rows[1:]:
            if len(r) < 3 or not r[0].isdigit(): continue
            pending_factor["levels"].append({
                "level": int(r[0]),
                "description": re.sub(r'\s+', ' ', r[1]),
                "points": int(r[2]),
            })
        factors.append(pending_factor); pending_factor = None

# ---------- 2. parse the xlsx Factor Breakdown (short labels + "what does this mean") ----------
wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
fb = wb['Factor Breakdown']
def g(r, c):
    v = fb.cell(r, c).value
    return re.sub(r'\s+', ' ', str(v).strip()) if v is not None else ""

# (factor name in docx, header row of the factor block, column pair start)
BLOCKS = [
    ("Professional Expertise",                      [(7,1),(7,3),(10,1),(10,3),(13,1),(13,3),(16,1),(16,3)]),
    ("Commercial and Operational Know How",         [(7,5),(7,7),(10,5),(10,7),(13,5),(13,7),(16,5)]),
    ("Working With and Through People to Deliver",  [(27,1),(27,3),(30,1),(30,3),(33,1),(33,3)]),
    ("Communicating",                               [(27,5),(27,7),(30,5),(30,7),(33,5),(33,7)]),
    ("Building Relationships",                      [(41,1),(41,3),(44,1),(44,3),(47,1),(47,3),(50,1)]),
    ("Analysis & Insight",                          [(61,1),(61,3),(64,1),(64,3),(67,1),(67,3)]),
    ("Developing Solutions",                        [(61,5),(61,7),(64,5),(64,7),(67,5),(67,7)]),
    ("Change & Innovation",                         [(75,1),(75,3),(78,1),(78,3),(81,1),(81,3)]),
    ("Responsibility & Ownership",                  [(92,1),(92,3),(95,1),(95,3),(98,1),(98,3),(101,1)]),
    ("Decision-Making",                             [(92,5),(92,7),(95,5),(95,7),(98,5),(98,7),(101,5)]),
    ("Context & Impact",                            [(109,1),(109,3),(112,1),(112,3),(115,1),(115,3)]),
]
extra = {}
for fname, cells in BLOCKS:
    extra[fname] = [{"label": g(r, c), "meaning": g(r+2, c)} for (r, c) in cells]

# ---------- 3. merge + validate ----------
byname = {f["name"]: f for f in factors}
assert len(factors) == 11, len(factors)
for fname, lst in extra.items():
    f = byname[fname]
    assert len(f["levels"]) == len(lst), (fname, len(f["levels"]), len(lst))
    for lv, ex in zip(f["levels"], lst):
        assert ex["label"] and ex["meaning"], (fname, lv["level"], ex)
        lv["label"] = ex["label"]
        lv["meaning"] = ex["meaning"]

# cross-check points against the Points sheet
pts = wb['Points']
COLS = {"Professional Expertise":3, "Commercial and Operational Know How":4,
        "Working With and Through People to Deliver":5, "Communicating":6,
        "Building Relationships":7, "Analysis & Insight":8, "Developing Solutions":9,
        "Change & Innovation":10, "Responsibility & Ownership":11,
        "Decision-Making":12, "Context & Impact":13}
for fname, col in COLS.items():
    for lv in byname[fname]["levels"]:
        cell = pts.cell(2 + lv["level"], col).value
        assert cell == lv["points"], (fname, lv["level"], cell, lv["points"])

print("validated:", len(factors), "factors,", sum(len(f['levels']) for f in factors), "levels")
print("max points:", sum(max(l['points'] for l in f['levels']) for f in factors))
print("min points:", sum(min(l['points'] for l in f['levels']) for f in factors))

# ---------- 4. emit TypeScript ----------
def slug(s):
    s = s.lower().replace("&", "and")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")

def q(s):
    return json.dumps(s, ensure_ascii=False)

CS_KEY = {"Know How": "knowHow", "People": "people", "Thinking": "thinking", "Delivery": "delivery"}
SHORT = {
    "Professional Expertise": "Professional Expertise",
    "Commercial and Operational Know How": "Commercial & Operational Know-How",
    "Working With and Through People to Deliver": "Working Through People",
    "Communicating": "Communicating",
    "Building Relationships": "Building Relationships",
    "Analysis & Insight": "Analysis & Insight",
    "Developing Solutions": "Developing Solutions",
    "Change & Innovation": "Change & Innovation",
    "Responsibility & Ownership": "Responsibility & Ownership",
    "Decision-Making": "Decision-Making",
    "Context & Impact": "Context & Impact",
}

HEADER = '''// AUTO-GENERATED from "Our role sizing scheme v2.docx" and "JE Scheme Database 2022.xlsx".
// Every level's points value is cross-validated against the workbook's Points sheet.
// Regenerate with scripts/generate-scheme.py rather than editing by hand.

export type CornerstoneKey = 'knowHow' | 'people' | 'thinking' | 'delivery'

export interface SchemeLevel {
  /** Level number as it appears in the scheme (1-based). */
  level: number
  /** Short headline for the level, from the Factor Breakdown sheet. */
  label: string
  /** Full scheme criteria, from the role sizing scheme document. */
  description: string
  /** The "what does this mean?" expansion from the Factor Breakdown sheet. */
  meaning: string
  points: number
}

export interface SchemeFactor {
  id: string
  /** Factor number 1-11, as numbered in the scheme document. */
  number: number
  name: string
  /** Column heading used in the JE Scheme Database workbook. */
  shortName: string
  cornerstone: CornerstoneKey
  levels: SchemeLevel[]
}

export interface Cornerstone {
  key: CornerstoneKey
  number: number
  name: string
  intro: string
}
'''

FOOTER = '''export const FACTORS_BY_ID: Record<string, SchemeFactor> = Object.fromEntries(
  FACTORS.map((f) => [f.id, f]),
)

/** Highest score the scheme can produce (1645). */
export const MAX_POINTS = FACTORS.reduce(
  (t, f) => t + Math.max(...f.levels.map((l) => l.points)),
  0,
)

/** Lowest score the scheme can produce (240) - every factor must be scored. */
export const MIN_POINTS = FACTORS.reduce(
  (t, f) => t + Math.min(...f.levels.map((l) => l.points)),
  0,
)

/** Maximum points available within each cornerstone, for weighting displays. */
export const CORNERSTONE_MAX: Record<CornerstoneKey, number> = CORNERSTONES.reduce(
  (acc, c) => {
    acc[c.key] = FACTORS.filter((f) => f.cornerstone === c.key).reduce(
      (t, f) => t + Math.max(...f.levels.map((l) => l.points)),
      0,
    )
    return acc
  },
  {} as Record<CornerstoneKey, number>,
)
'''

out = [HEADER, "export const CORNERSTONES: Cornerstone[] = ["]
for c in cornerstones:
    out += ["  {", f"    key: '{CS_KEY[c['name']]}',", f"    number: {c['num']},",
            f"    name: {q(c['name'])},", f"    intro:\n      {q(c['intro'])},", "  },"]
out += ["]\n", "export const FACTORS: SchemeFactor[] = ["]
for f in factors:
    out += ["  {", f"    id: '{slug(f['name'])}',", f"    number: {f['num']},",
            f"    name: {q(f['name'])},", f"    shortName: {q(SHORT[f['name']])},",
            f"    cornerstone: '{CS_KEY[f['cornerstone']]}',", "    levels: ["]
    for l in f["levels"]:
        out += ["      {", f"        level: {l['level']},", f"        label: {q(l['label'])},",
                f"        description:\n          {q(l['description'])},",
                f"        meaning:\n          {q(l['meaning'])},",
                f"        points: {l['points']},", "      },"]
    out += ["    ],", "  },"]
out += ["]\n", FOOTER]

dest = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src", "scheme", "scheme.ts")
with open(dest, "w") as fh:
    fh.write("\n".join(out))
print("wrote", os.path.normpath(dest))
